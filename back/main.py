import io
import json
import logging
import os
import re
import tarfile
import threading

import requests
from fastapi import BackgroundTasks, FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from langchain_ollama import ChatOllama
from openpyxl import Workbook
from openpyxl.styles import Alignment

from config import OLLAMA_BASE_URL, OLLAMA_MODEL

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

llm = ChatOllama(
    model=OLLAMA_MODEL,
    temperature=0,
    base_url=OLLAMA_BASE_URL,
)

# Estado da avaliacao em andamento
avaliacao_estado = {
    "rodando": False,
    "progresso": "",
    "concluida": False,
}
avaliacao_lock = threading.Lock()

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "avaliacao.xlsx")


def limpar_think_tags(texto: str) -> str:
    """Remove tags <think>...</think> que alguns modelos (qwen) retornam."""
    return re.sub(r"<think>.*?</think>", "", texto, flags=re.DOTALL).strip()


def recarregar_modelo():
    """Descarrega e recarrega o modelo no Ollama para limpar contexto."""
    try:
        # Descarregar
        requests.post(
            f"{OLLAMA_BASE_URL}/api/generate",
            json={"model": OLLAMA_MODEL, "keep_alive": 0},
            timeout=30,
        )
        # Recarregar (warm up com prompt vazio)
        requests.post(
            f"{OLLAMA_BASE_URL}/api/generate",
            json={"model": OLLAMA_MODEL, "prompt": "", "keep_alive": "10m"},
            timeout=60,
        )
        logger.info("Modelo recarregado no Ollama")
    except Exception as e:
        logger.warning(f"Falha ao recarregar modelo: {e}")


def filtrar_arquivos(arquivos: dict[str, str], filtro: str) -> dict[str, str]:
    """Filtra arquivos por regex ou substring. Retorna os que batem."""
    if not filtro:
        return arquivos

    encontrados = {}
    for nome, conteudo in arquivos.items():
        try:
            if re.search(filtro, nome, re.IGNORECASE):
                encontrados[nome] = conteudo
        except re.error:
            if filtro.lower() in nome.lower():
                encontrados[nome] = conteudo
    return encontrados


def montar_codigo(arquivos_filtrados: dict[str, str]) -> tuple[str, str]:
    """
    Monta o codigo a partir dos arquivos filtrados.
    Retorna (codigo, nome_descritivo).
    """
    if len(arquivos_filtrados) == 1:
        nome = list(arquivos_filtrados.keys())[0]
        return arquivos_filtrados[nome], nome

    # Multiplos arquivos: concatenar com separadores
    partes = []
    for nome, conteudo in arquivos_filtrados.items():
        partes.append(f"=== ARQUIVO: {nome} ===\n{conteudo}")
    nomes = ", ".join(arquivos_filtrados.keys())
    return "\n\n".join(partes), nomes


def avaliar_codigo(codigo: str, nome_arquivo: str, pergunta: str, opcoes: list[str]) -> str:
    """
    Agente avaliador: recebe codigo e uma pergunta com opcoes.
    Retorna somente a opcao escolhida.
    """
    opcoes_formatadas = "\n".join(f"  - {op}" for op in opcoes)
    prompt = (
        f"Voce eh um AGENTE AVALIADOR de codigo. "
        f"RESPONDA EXCLUSIVAMENTE EM PORTUGUES.\n\n"
        f"REGRAS:\n"
        f"1. Responda SOMENTE com a opcao escolhida, copiada EXATAMENTE como esta listada.\n"
        f"2. NAO escreva explicacoes, exemplos de codigo, markdown, nem texto adicional.\n"
        f"3. Sua resposta deve conter APENAS a opcao escolhida e NADA mais.\n\n"
        f"ARQUIVO(S) SENDO ANALISADO(S): {nome_arquivo}\n"
        f"CODIGO:\n```\n{codigo}\n```\n\n"
        f"PERGUNTA: {pergunta}\n"
        f"OPCOES:\n{opcoes_formatadas}\n\n"
        f"SUA RESPOSTA (somente a opcao):"
    )
    logger.info(f"  LLM <- {pergunta[:60]}...")
    response = llm.invoke(prompt)
    raw = limpar_think_tags(response.content).strip('"').strip("'")

    # Pegar so a primeira linha caso o modelo escreva mais
    resposta = raw.split("\n")[0].strip().strip('"').strip("'")
    # Remover justificativa se o modelo adicionou com |
    if "|" in resposta:
        resposta = resposta.split("|")[0].strip().strip('"').strip("'")

    logger.info(f"  LLM -> {resposta}")
    return resposta


def justificar_resposta(codigo: str, nome_arquivo: str, pergunta: str, resposta: str) -> str:
    """Pede ao LLM uma justificativa curta para resposta negativa/parcial."""
    prompt = (
        f"Voce eh um AGENTE AVALIADOR de codigo. "
        f"RESPONDA EXCLUSIVAMENTE EM PORTUGUES.\n\n"
        f"O codigo abaixo foi avaliado com a seguinte pergunta e obteve a resposta indicada.\n"
        f"Explique de forma CURTA e DIRETA (1-2 frases) o motivo da resposta.\n\n"
        f"ARQUIVO(S): {nome_arquivo}\n"
        f"CODIGO:\n```\n{codigo}\n```\n\n"
        f"PERGUNTA: {pergunta}\n"
        f"RESPOSTA DADA: {resposta}\n\n"
        f"JUSTIFICATIVA CURTA:"
    )
    logger.info(f"  LLM justificativa <- {pergunta[:40]}...")
    response = llm.invoke(prompt)
    justificativa = limpar_think_tags(response.content).strip()
    # Pegar so primeiras linhas
    linhas = [l.strip() for l in justificativa.split("\n") if l.strip()]
    justificativa = " ".join(linhas[:3])
    logger.info(f"  LLM justificativa -> {justificativa[:80]}...")
    return justificativa


def resumir_comentarios(justificativas: list[str]) -> str:
    """Pede ao LLM para resumir todas as justificativas em um comentario unico."""
    todas = "\n".join(f"- {j}" for j in justificativas)
    prompt = (
        f"Voce eh um AGENTE AVALIADOR de codigo. "
        f"RESPONDA EXCLUSIVAMENTE EM PORTUGUES.\n\n"
        f"Abaixo estao os problemas encontrados na avaliacao de um aluno.\n"
        f"Faca um RESUMO CURTO e DIRETO de todos os problemas em um unico paragrafo.\n"
        f"NAO use markdown, bullets ou formatacao. Apenas texto corrido.\n\n"
        f"PROBLEMAS:\n{todas}\n\n"
        f"RESUMO:"
    )
    logger.info(f"  LLM resumo comentarios <- {len(justificativas)} justificativas")
    response = llm.invoke(prompt)
    resumo = limpar_think_tags(response.content).strip()
    # Limpar formatacao extra
    linhas = [l.strip() for l in resumo.split("\n") if l.strip()]
    resumo = " ".join(linhas)
    logger.info(f"  LLM resumo -> {resumo[:100]}...")
    return resumo


def extrair_tar(tar_bytes: bytes, filename: str) -> dict[str, str]:
    """Extrai o tar.gz UMA vez e retorna {nome_arquivo: conteudo}."""
    arquivos = {}
    with tarfile.open(fileobj=io.BytesIO(tar_bytes), mode="r:gz") as tf:
        for member in tf.getmembers():
            if not member.isfile():
                continue
            nome_arquivo = os.path.basename(member.name)
            if not nome_arquivo or nome_arquivo.startswith("."):
                continue
            try:
                f = tf.extractfile(member)
                if f is None:
                    continue
                conteudo = f.read().decode("utf-8", errors="replace")
                arquivos[nome_arquivo] = conteudo
            except Exception:
                continue
    logger.info(f"[{filename}] {len(arquivos)} arquivos extraidos: {list(arquivos.keys())}")
    return arquivos


def processar_avaliacao(
    perguntas: list[dict],
    todos_tar: list[tuple[str, bytes]],
):
    """Executa a avaliacao em background, salvando o Excel incrementalmente."""
    with avaliacao_lock:
        avaliacao_estado["rodando"] = True
        avaliacao_estado["concluida"] = False
        avaliacao_estado["progresso"] = "Iniciando..."

    wb = Workbook()
    ws = wb.active
    ws.title = "Avaliacao"
    wrap_align = Alignment(wrap_text=True, vertical="top")

    # Cabecalho: Matricula | Pergunta1 | Pergunta2 | ... | Comentários
    ws.cell(row=1, column=1, value="Matricula").alignment = wrap_align
    for i, p in enumerate(perguntas):
        # Quebrar texto da pergunta com newlines para ficar compacto
        texto_header = p["pergunta"].replace(" ", "\n", p["pergunta"].count(" ") // 2) if len(p["pergunta"]) > 30 else p["pergunta"]
        cell = ws.cell(row=1, column=2 + i, value=texto_header)
        cell.alignment = wrap_align
    col_comentarios = 2 + len(perguntas)
    ws.cell(row=1, column=col_comentarios, value="Comentários").alignment = wrap_align

    row = 2
    total_tars = len(todos_tar)

    for t_idx, (nome_tar, tar_bytes) in enumerate(todos_tar):
        matricula = nome_tar.replace(".tar.gz", "").replace(".tgz", "")
        logger.info(f"=== [{t_idx+1}/{total_tars}] Matricula: {matricula} ===")

        # Recarregar modelo a cada tar.gz para limpar contexto
        with avaliacao_lock:
            avaliacao_estado["progresso"] = f"Recarregando modelo para {matricula} ({t_idx+1}/{total_tars})"
        recarregar_modelo()

        with avaliacao_lock:
            avaliacao_estado["progresso"] = f"Extraindo {matricula} ({t_idx+1}/{total_tars})"

        arquivos = extrair_tar(tar_bytes, nome_tar)

        ws.cell(row=row, column=1, value=matricula)

        # Acumular justificativas para respostas negativas/parciais
        justificativas_acumuladas: list[str] = []

        # Para cada pergunta, buscar o arquivo correspondente
        for i, p in enumerate(perguntas):
            filtro = p.get("arquivo", "").strip()
            with avaliacao_lock:
                avaliacao_estado["progresso"] = f"{matricula} - pergunta {i+1}/{len(perguntas)} ({t_idx+1}/{total_tars})"

            # Filtrar arquivos pelo regex da pergunta
            arquivos_filtrados = filtrar_arquivos(arquivos, filtro)

            if not arquivos_filtrados:
                ws.cell(row=row, column=2 + i, value="N/A")
                logger.info(f"  [{matricula}] {p['pergunta'][:40]}... -> N/A (nenhum arquivo bate com '{filtro}')")
                continue

            codigo, nome_desc = montar_codigo(arquivos_filtrados)
            resposta = avaliar_codigo(codigo, nome_desc, p["pergunta"], p["opcoes"])

            # Marcar se mais de um arquivo bateu
            if len(arquivos_filtrados) > 1:
                resposta = f"[MULTI-ARQUIVO] {resposta}"
                logger.info(f"  [{matricula}] {p['pergunta'][:40]}... -> {resposta} (arquivos: {list(arquivos_filtrados.keys())})")
            else:
                logger.info(f"  [{matricula}] {p['pergunta'][:40]}... -> {resposta}")

            ws.cell(row=row, column=2 + i, value=resposta)

            # Se resposta nao eh "Sim", pedir justificativa
            resposta_limpa = resposta.replace("[MULTI-ARQUIVO] ", "")
            if resposta_limpa.strip().upper() != "SIM":
                try:
                    justificativa = justificar_resposta(codigo, nome_desc, p["pergunta"], resposta_limpa)
                    justificativas_acumuladas.append(f"{p['pergunta'][:50]}: {justificativa}")
                except Exception as e:
                    logger.warning(f"  Erro ao justificar: {e}")

        # Gerar comentario resumido se houve respostas negativas/parciais
        if justificativas_acumuladas:
            with avaliacao_lock:
                avaliacao_estado["progresso"] = f"{matricula} - gerando comentário ({t_idx+1}/{total_tars})"
            try:
                comentario = resumir_comentarios(justificativas_acumuladas)
            except Exception as e:
                logger.warning(f"  Erro ao resumir comentarios: {e}")
                comentario = "; ".join(justificativas_acumuladas)
            cell = ws.cell(row=row, column=col_comentarios, value=comentario)
            cell.alignment = wrap_align
            logger.info(f"  [{matricula}] Comentário: {comentario[:80]}...")

        row += 1

        # Salvar Excel parcialmente apos cada matricula
        wb.save(EXCEL_PATH)
        logger.info(f"  Excel salvo ({row - 2} matriculas)")

    with avaliacao_lock:
        avaliacao_estado["rodando"] = False
        avaliacao_estado["concluida"] = True
        avaliacao_estado["progresso"] = f"Concluida! {row - 2} matriculas avaliadas."

    logger.info(f"Avaliacao 100% concluida! {row - 2} matriculas no Excel.")


@app.post("/avaliar")
async def avaliar(
    questions: str = Form(...),
    tar_files: list[UploadFile] = File(...),
    background_tasks: BackgroundTasks = BackgroundTasks(),
):
    """Recebe perguntas + tar.gz e inicia avaliacao em background."""
    with avaliacao_lock:
        if avaliacao_estado["rodando"]:
            return JSONResponse(
                status_code=409,
                content={"error": "Ja existe uma avaliacao em andamento."},
            )

    perguntas = json.loads(questions)
    logger.info(f"Recebido {len(perguntas)} perguntas, {len(tar_files)} arquivo(s)")
    for i, p in enumerate(perguntas):
        logger.info(f"  Pergunta {i+1}: {p['pergunta']} | Arquivo: {p.get('arquivo', '')} | Opcoes: {p['opcoes']}")

    # Ler todos os tar.gz em memoria antes de ir para background
    todos_tar: list[tuple[str, bytes]] = []
    for tf in tar_files:
        tar_bytes = await tf.read()
        nome = tf.filename or "desconhecido"
        todos_tar.append((nome, tar_bytes))
        logger.info(f"  Arquivo recebido: {nome} ({len(tar_bytes)} bytes)")

    # Remover Excel anterior se existir
    if os.path.exists(EXCEL_PATH):
        os.remove(EXCEL_PATH)

    background_tasks.add_task(processar_avaliacao, perguntas, todos_tar)

    return JSONResponse(content={"status": "iniciado", "total_arquivos": len(todos_tar)})


@app.get("/status")
def get_status():
    """Retorna o estado atual da avaliacao."""
    with avaliacao_lock:
        return JSONResponse(content={
            "rodando": avaliacao_estado["rodando"],
            "concluida": avaliacao_estado["concluida"],
            "progresso": avaliacao_estado["progresso"],
        })


@app.get("/download")
def download():
    """Baixa o Excel (parcial ou completo)."""
    if not os.path.exists(EXCEL_PATH):
        return JSONResponse(status_code=404, content={"error": "Nenhuma avaliacao disponivel."})
    return FileResponse(
        EXCEL_PATH,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="avaliacao.xlsx",
    )
